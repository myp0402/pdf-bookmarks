#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成测试样例：一份 docx + 一份模仿 Mac 版 Word 导出结果的 PDF（无书签）。

刻意埋了这几个坑，用来验证 bookmark_pdf.py 的鲁棒性：
  1. 目录页：标题文字在第 1 页以“点线+页码”形式出现，不能误当成正文标题
  2. 页眉页脚：每页都有页眉文字，不能参与匹配
  3. 自动编号：docx 标题不含“第一章 / 1.1”，PDF 里带，必须靠后缀匹配
  4. 正文误导：正文里出现“目标与范围”四字，且与真标题同页
  5. 标题折行：长标题在 PDF 里被折成两行
"""
import os
from docx import Document
from docx.shared import Pt
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

HERE = os.path.dirname(os.path.abspath(__file__))
FONT = "/System/Library/Fonts/Supplemental/Arial Unicode.ttf"

pdfmetrics.registerFont(TTFont("CJK", FONT))

W, H = 595, 842          # A4
LM, RM = 72, 72
TOP = H - 90

HEADERS = [
    ("第一章 引言", 18),
    ("1.1 背景", 14),
]
BODY = "云鼎科技研究院在 2025 年完成了全部既定实验，相关数据已归档。"


def header_footer(c, pno, total):
    c.setFont("CJK", 9)
    c.setFillGray(0.45)
    c.drawString(LM, H - 42, "云鼎科技 2025 年度技术报告")          # 页眉
    c.drawRightString(W - RM, 40, f"第 {pno} 页 共 {total} 页")      # 页脚
    c.setFillGray(0)


def heading(c, y, text, size):
    c.setFont("CJK", size)
    c.drawString(LM, y, text)
    return y - size - 12


def para(c, y, text, size=11):
    c.setFont("CJK", size)
    c.setFillGray(0.15)
    c.drawString(LM, y, text)
    c.setFillGray(0)
    return y - size - 14


def build_pdf(path, bare_toc=False, toc_size=11):
    total = 6
    c = canvas.Canvas(path, pagesize=(W, H))
    c.setTitle("云鼎科技 2025 年度技术报告")

    # ---------------- p1 目录
    header_footer(c, 1, total)
    c.setFont("CJK", 16)
    c.drawString(LM, TOP, "目录")
    y = TOP - 34
    toc = [
        ("第一章 引言", 2), ("1.1 背景", 2), ("1.2 目标与范围", 3),
        ("第二章 方法与实现", 4), ("2.1 数据来源", 4),
        ("2.2 一个特别长的标题，用来测试 PDF 里的折行情况", 5),
        ("第三章 结论", 6),
    ]
    for t, p in toc:
        c.setFont("CJK", toc_size)
        if bare_toc:
            c.drawString(LM, y, t)                    # 无点线、无页码的目录
        else:
            dots = "." * max(4, int((440 - len(t) * 11) / 5))
            c.drawString(LM, y, f"{t} {dots} {p}")
        y -= 24
    c.showPage()

    # ---------------- p2
    header_footer(c, 2, total)
    y = TOP
    y = heading(c, y, "第一章 引言", 18)
    y = para(c, y, BODY)
    y = para(c, y, "本节交代研究背景与边界。")
    y -= 10
    y = heading(c, y, "1.1 背景", 14)
    y = para(c, y, BODY)
    y = para(c, y, "行业在过去三年经历了显著变化。")
    c.showPage()

    # ---------------- p3  正文误导 + 真标题同页
    header_footer(c, 3, total)
    y = TOP
    y = para(c, y, "关于目标与范围的说明见下文。")     # ← 陷阱
    y -= 10
    y = heading(c, y, "1.2 目标与范围", 14)
    y = para(c, y, BODY)
    c.showPage()

    # ---------------- p4
    header_footer(c, 4, total)
    y = TOP
    y = heading(c, y, "第二章 方法与实现", 18)
    y = para(c, y, BODY)
    y -= 10
    y = heading(c, y, "2.1 数据来源", 14)
    y = para(c, y, BODY)
    c.showPage()

    # ---------------- p5  折行标题
    header_footer(c, 5, total)
    y = TOP
    c.setFont("CJK", 14)
    c.drawString(LM, y, "2.2 一个特别长的标题，用来测试")     # ← 折行上半
    c.drawString(LM, y - 20, "PDF 里的折行情况")              # ← 折行下半
    y -= 44
    y = para(c, y, BODY)
    c.showPage()

    # ---------------- p6
    header_footer(c, 6, total)
    y = TOP
    y = heading(c, y, "第三章 结论", 18)
    y = para(c, y, BODY)
    c.showPage()

    c.save()


def build_docx(path):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)

    def add(text, lvl):
        doc.add_heading(text, level=lvl)

    add("引言", 1)
    doc.add_paragraph(BODY)
    add("背景", 2)
    doc.add_paragraph(BODY)
    add("目标与范围", 2)
    doc.add_paragraph(BODY)
    add("方法与实现", 1)
    doc.add_paragraph(BODY)
    add("数据来源", 2)
    doc.add_paragraph(BODY)
    add("一个特别长的标题，用来测试 PDF 里的折行情况", 2)   # 折行
    doc.add_paragraph(BODY)
    add("结论", 1)
    doc.add_paragraph(BODY)
    doc.save(path)


if __name__ == "__main__":
    build_docx(os.path.join(HERE, "fixture.docx"))
    build_pdf(os.path.join(HERE, "fixture.pdf"))

    print("已生成 fixture.docx / fixture.pdf")
